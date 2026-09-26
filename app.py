import io
import os
import secrets
import tempfile
import re
import ast
import json
import time
import base64
import hashlib
import concurrent.futures
from pathlib import Path
from urllib.parse import urlparse
import requests
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from dotenv import load_dotenv
from flask import Flask, request, jsonify, send_from_directory, send_file, g
from werkzeug.utils import secure_filename
from werkzeug.exceptions import HTTPException
from pypdf import PdfReader
from core.llm.client import generate, chat
from core.rag.retrieve import retrieve
from core.vision.image_handler import transcribe
from core.vision.step_checker import check_working
import core.misconceptions
from diagram_shots import shot_messages
from sessionlog import log_turn, attachment_meta, list_sessions, session_turns, all_sessions, session_turns_admin, stats
import pdftext
import users
import quizbank
import settings as _settings
from api_v1 import v1
from ratelimit import RateLimiter, SlidingWindowLimiter
from core.prompts import build_check_prompt, tutor_reply as _tutor_reply, TUTOR_SYSTEM_PROMPT, hint_note
load_dotenv()
app = Flask(__name__, static_folder='static')
users.init()
app.register_blueprint(v1)
users.purge_expired()

@app.errorhandler(Exception)
def _unhandled_exception(e):
    if isinstance(e, HTTPException):
        return (jsonify({'error': e.description}), e.code)
    app.logger.exception(e)
    if isinstance(e, ConnectionError):
        return (jsonify({'error': 'Could not reach the model. Is Ollama running?'}), 503)
    try:
        from groq import APITimeoutError as _GroqTimeout, RateLimitError as _GroqRate
        if isinstance(e, _GroqTimeout):
            return (jsonify({'error': 'The model took too long. Try again.'}), 503)
        if isinstance(e, _GroqRate):
            return (jsonify({'error': 'Too many requests. Slow down a bit and try again.'}), 429)
    except ImportError:
        pass
    return (jsonify({'error': 'Something went wrong on our end. My bad gng'}), 500)

@app.errorhandler(505)
def _http_version_not_supported(e):
    return (jsonify({'error': 'HTTP version not supported.THIS IS NOT OUR FAULT'}), 505)
SESSION_COOKIE = 'ollm_session'
COOKIE_SECURE = os.environ.get('OLLM_HTTPS', '').strip().lower() in {'1', 'true', 'yes'}
_login_limiter = RateLimiter(max_failures=8, window_seconds=15 * 60)
_INGEST_STATUS: dict = {}
_llm_limiter = SlidingWindowLimiter(max_calls=20, window_seconds=60)

def _check_llm_rate() -> object:
    if not _settings.load()['rate_limiting']:
        return None
    key = f"user:{g.user['id']}"
    wait = _llm_limiter.check_and_record(key)
    if wait:
        resp = jsonify({'error': 'Too many requests. Slow down and try again shortly.'})
        resp.headers['Retry-After'] = str(wait)
        return (resp, 429)
    return None

def _login_keys(username: str) -> tuple:
    return (f'ip:{request.remote_addr}', f"user:{(username or '').strip().lower()}")
_OPEN_ENDPOINTS = {'index', 'static', 'api_register', 'api_login', 'api_logout', 'api_me', 'admin_panel', 'teacher_panel'}
_DEV_USER = {'id': 1, 'username': 'dev', 'role': 'admin'}

@app.before_request
def _load_user():
    cfg = _settings.load()
    dev_env = os.environ.get('DEV_BYPASS_AUTH', '').lower() in {'1', 'true', 'yes'}
    if dev_env or not cfg['auth_required']:
        g.user = _DEV_USER
        return None
    g.user = users.user_for_token(request.cookies.get(SESSION_COOKIE, ''))
    if request.blueprint == 'v1':
        return None
    if request.endpoint in _OPEN_ENDPOINTS:
        return None
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    return None

def _auth_response(user: dict, token: str):
    resp = jsonify({'user': {'id': user['id'], 'username': user['username']}})
    resp.set_cookie(SESSION_COOKIE, token, max_age=users.SESSION_TTL_SECONDS, httponly=True, samesite='Lax', secure=COOKIE_SECURE)
    return resp

@app.route('/api/register', methods=['POST'])
def api_register():
    data = request.json or {}
    ip_key = f'register-ip:{request.remote_addr}'
    wait = _login_limiter.retry_after(ip_key)
    if wait:
        resp = jsonify({'error': 'Too many attempts. Try again later.'})
        resp.headers['Retry-After'] = str(wait)
        return (resp, 429)
    try:
        uid = users.create_user(data.get('username', ''), data.get('password', ''))
    except users.UserError as e:
        _login_limiter.record_failure(ip_key)
        return (jsonify({'error': str(e)}), 400)
    _login_limiter.clear(ip_key)
    username = (data.get('username') or '').strip()
    return _auth_response({'id': uid, 'username': username}, users.create_session(uid))

@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.json or {}
    username = (data.get('username') or '').strip()
    keys = _login_keys(username)
    wait = _login_limiter.retry_after(*keys)
    if wait:
        resp = jsonify({'error': 'Too many failed logins. Try again later.'})
        resp.headers['Retry-After'] = str(wait)
        return (resp, 429)
    uid = users.authenticate(data.get('username', ''), data.get('password', ''))
    if uid is None:
        _login_limiter.record_failure(*keys)
        return (jsonify({'error': 'Wrong username or password.'}), 401)
    _login_limiter.clear(*keys)
    users.purge_expired()
    return _auth_response({'id': uid, 'username': username}, users.create_session(uid))

@app.route('/api/logout', methods=['POST'])
def api_logout():
    users.delete_session(request.cookies.get(SESSION_COOKIE, ''))
    resp = jsonify({'ok': True})
    resp.delete_cookie(SESSION_COOKIE)
    return resp

@app.route('/api/me')
def api_me():
    return jsonify({'user': g.user})

def _key(data: dict, field: str, env_var: str) -> str:
    return (data.get(field) or '').strip() or os.environ.get(env_var, '').strip()

def _curriculum_params(data: dict) -> tuple[str, str]:
    cid = (data.get('curriculum_id') or '').strip()
    if cid:
        try:
            manifest = _cur.load(cid)
            from core.prompts import build_tutor_prompt
            prompt = build_tutor_prompt(flavour=manifest['prompt_flavour'])
            collection = manifest.get('collection', 'reference')
            return (prompt, collection)
        except Exception:
            pass
    return (TUTOR_SYSTEM_PROMPT, 'reference')
GRAPH_SYSTEM_PROMPT = "You are a Python matplotlib code generator. Output ONLY valid Python code with no explanation and no markdown fences. The following are already in scope — do NOT import or redefine them: plt, np. Write only the plotting logic. Rules:\n- Use LaTeX notation in all labels, titles, and legends: wrap math in $...$. Example: label=r'$f(x) = x^{2}$'. Always use r-strings for LaTeX.\n- Do NOT call plt.savefig(), plt.show(), plt.close(), or import anything.\n- End with plt.tight_layout().\n- NEVER invent real-world data. If the graph needs figures you do not have — prices, populations, sales, historical values — do not guess them. Output exactly NEED_REAL_DATA and nothing else.\n- Every worked example uses ONE standard market: demand P = 10 - 0.8Q, supply P = 2 + 0.8Q. They cross at Q = 5, P = 6. Reuse those equations.\n- A ceiling or maximum must sit BELOW P = 6. A floor, minimum or union wage must sit ABOVE P = 6. A line level with 6 shows nothing.\n- A horizontal line at price p meets demand at Q = (10 - p) / 0.8 and supply at Q = (p - 2) / 0.8. Compute both before placing any marker.\n- Never put a marker at a point that is not on a curve you plotted."
_GRAPH_ALLOWED_NODES = (ast.Module, ast.Expr, ast.Assign, ast.AugAssign, ast.For, ast.If, ast.Pass, ast.Call, ast.keyword, ast.Name, ast.Attribute, ast.Subscript, ast.Slice, ast.Starred, ast.Constant, ast.Tuple, ast.List, ast.Dict, ast.Set, ast.ListComp, ast.SetComp, ast.DictComp, ast.GeneratorExp, ast.comprehension, ast.BinOp, ast.UnaryOp, ast.BoolOp, ast.Compare, ast.IfExp, ast.operator, ast.unaryop, ast.boolop, ast.cmpop, ast.expr_context)
_GRAPH_BANNED_ATTRS = {'savefig', 'show', 'close', 'imsave'}
_GRAPH_BANNED_NAMES = {'exec', 'eval', 'compile', 'open', 'input', 'breakpoint', 'help', 'globals', 'locals', 'vars', 'dir', 'getattr', 'setattr', 'delattr'}
_GRAPH_BUILTINS = {'abs': abs, 'min': min, 'max': max, 'sum': sum, 'round': round, 'pow': pow, 'range': range, 'len': len, 'enumerate': enumerate, 'zip': zip, 'sorted': sorted, 'reversed': reversed, 'int': int, 'float': float, 'str': str, 'bool': bool, 'list': list, 'tuple': tuple, 'dict': dict, 'set': set}

def _check_graph_code(code: str) -> str | None:
    try:
        tree = ast.parse(code)
    except SyntaxError as e:
        return f'it did not parse: {e.msg}'
    for node in ast.walk(tree):
        if not isinstance(node, _GRAPH_ALLOWED_NODES):
            return f'{type(node).__name__} is not allowed'
        if isinstance(node, ast.Attribute):
            if node.attr.startswith('_'):
                return f'attribute {node.attr!r} is not allowed'
            if node.attr in _GRAPH_BANNED_ATTRS:
                return f'{node.attr}() is not allowed — the server renders the figure'
        if isinstance(node, ast.Name):
            if node.id.startswith('_'):
                return f'name {node.id!r} is not allowed'
            if node.id in _GRAPH_BANNED_NAMES:
                return f'{node.id}() is not allowed'
    return None

def _render_png() -> str:
    buf = io.BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight', dpi=150)
    plt.close('all')
    buf.seek(0)
    return base64.b64encode(buf.read()).decode('utf-8')
_LITERAL_DATA_MIN = 4

def _has_literal_data(code: str) -> bool:
    try:
        tree = ast.parse(code)
    except SyntaxError:
        return False
    for node in ast.walk(tree):
        if isinstance(node, (ast.List, ast.Tuple)):
            numbers = [e for e in node.elts if isinstance(e, ast.Constant) and isinstance(e.value, (int, float)) and (not isinstance(e.value, bool))]
            if len(numbers) >= _LITERAL_DATA_MIN:
                return True
    return False
ALLOWED_IMAGES = {'.png', '.jpg', '.jpeg', '.webp', '.gif'}
PDF_TEXT_LIMIT = 8000
PAPER_CACHE_DIR = Path(__file__).parent / 'data' / 'papers'
PAPER_MAX_BYTES = 20 * 1024 * 1024
PAPER_MAX_FETCH = 3
VERIFY_OVERSHOOT = 5
_PAPER_ID_RE = re.compile('^[0-9a-f]{16}$')
_GREETING_RE = re.compile('^\\s*(?:h(?:i|ey|ello|igh)|yo|sup|good\\s+(?:morning|afternoon|evening))(?:\\s+there)?\\s*[!.?]*\\s*$', re.IGNORECASE)
_FILLER_WORDS = {'test', 'testing', 'asdf', 'abc', 'stuff', 'thing', 'things', 'anything', 'something', 'idk', 'none', 'na'}

def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return '\n'.join(pdftext.doc_pages(reader))
_QUESTION_NUM_RE = re.compile('^\\s*\\d{1,2}[\\.\\)]\\s+\\S', re.MULTILINE)
_MATH_CHARS = set('=+×÷≤≥∫√^')

def _question_pages(data: bytes, budget: int=7000) -> tuple[bool, str]:
    reader = PdfReader(io.BytesIO(data))
    has_text = False
    scored = []
    for i, text in enumerate(pdftext.doc_pages(reader)):
        if text.strip():
            has_text = True
        if len(text.strip()) < 80:
            continue
        score = len(_QUESTION_NUM_RE.findall(text)) * 3 + text.count('?') + sum((c in _MATH_CHARS for c in text)) // 4
        if score >= 3:
            scored.append((score, i, text))
    if not scored:
        return (has_text, '')
    scored.sort(key=lambda t: -t[0])
    kept, total = ([], 0)
    for score, i, text in scored:
        if total >= budget:
            break
        kept.append((i, text))
        total += len(text)
    kept.sort()
    return (has_text, '\n\n'.join((text for _, text in kept))[:budget])

@app.route('/')
def index():
    resp = send_from_directory('static', 'index.html')
    resp.headers['Cache-Control'] = 'no-store'
    return resp

@app.route('/favicon.ico')
def favicon():
    return ('', 204)

@app.route('/api/upload', methods=['POST'])
def api_upload():
    file = request.files.get('file')
    if not file:
        return (jsonify({'error': 'No file provided'}), 400)
    name = file.filename.lower()
    ext = '.' + name.rsplit('.', 1)[-1] if '.' in name else ''
    if ext == '.pdf':
        text = _pdf_text(file.read())
        return jsonify({'type': 'pdf', 'text': text[:PDF_TEXT_LIMIT], 'name': file.filename})
    if ext in ALLOWED_IMAGES:
        b64 = base64.b64encode(file.read()).decode('utf-8')
        return jsonify({'type': 'image', 'b64': b64, 'name': file.filename})
    return (jsonify({'error': f'Unsupported file type: {ext}'}), 400)

@app.route('/api/graph', methods=['POST'])
def api_graph():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json
    description = data.get('description', '')
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    if not description:
        return (jsonify({'error': 'No description provided'}), 400)
    messages = [{'role': 'system', 'content': GRAPH_SYSTEM_PROMPT}]
    messages += shot_messages(description)
    messages.append({'role': 'user', 'content': f'Generate a graph for: {description}'})
    raw_code = chat(messages, model=model, host=host, provider=provider, api_key=api_key)
    if 'NEED_REAL_DATA' in raw_code:
        return (jsonify({'error': 'That graph needs real data I do not have. Give me the figures and I will plot them, or ask for a diagram that shows the shape of the relationship instead.'}), 400)
    code = re.sub('^```(?:python)?\\n?', '', raw_code.strip())
    code = re.sub('\\n?```$', '', code)
    code = re.sub('plt\\.(savefig|show|close)\\(.*?\\)\\n?', '', code)
    code = re.sub('^\\s*(?:import|from)\\s+.*$', '', code, flags=re.MULTILINE)
    err = _check_graph_code(code)
    if err:
        return (jsonify({'error': f'Refused to run the generated graph code — {err}'}), 400)
    try:
        plt.close('all')
        exec(code, {'__builtins__': _GRAPH_BUILTINS, 'plt': plt, 'np': np})
        if _has_literal_data(code):
            plt.gcf().text(0.99, 0.01, 'Illustrative — values are not real data', ha='right', va='bottom', fontsize=7, color='#b45309')
        return jsonify({'image': _render_png()})
    except Exception as e:
        plt.close('all')
        return (jsonify({'error': str(e)}), 500)

@app.route('/api/start', methods=['POST'])
def api_start():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json
    topic = data.get('topic', '')
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    attachment = data.get('attachment')
    vision_model = data.get('vision_model', '')
    system_prompt, collection = _curriculum_params(data)
    cid = (data.get('curriculum_id') or '').strip()
    if cid:
        q_path = Path('curricula') / cid / 'questions.json'
        if q_path.exists():
            try:
                qs = json.loads(q_path.read_text(encoding='utf-8'))
                if qs:
                    q_lines = []
                    for q in qs[:20]:
                        line = '- ' + (q.get('text') or q.get('question') or '').strip()
                        if q.get('answer'):
                            line += f" [answer: {q['answer'].strip()}]"
                        if line.strip() != '-':
                            q_lines.append(line)
                    if q_lines:
                        system_prompt += '\n\nQuestion bank for this curriculum (use these when quizzing the student — present them one at a time, never reveal the answer until the student has attempted it):\n' + '\n'.join(q_lines)
            except Exception:
                pass
    context_chunks = retrieve(topic, collection_name=collection)
    context = '\n\n'.join(context_chunks) if context_chunks else ''
    context_prefix = f'Reference material:\n{context}\n\n' if context else ''
    user_msg = {'role': 'user', 'content': f'{context_prefix}{topic}'}
    if attachment:
        _apply_attachment(user_msg, attachment, provider, model, host, api_key, vision_model)
    messages = [{'role': 'system', 'content': system_prompt}, user_msg]
    started = time.monotonic()
    raw = chat(messages, model=model, host=host, provider=provider, api_key=api_key)
    response, _ = _tutor_reply(raw)
    messages.append({'role': 'assistant', 'content': response})
    log_turn('start', session_id=data.get('session_id', ''), user_id=g.user['id'], mode='learn', topic=topic, provider=provider, model=model, input=topic, output=response, attachment=attachment_meta(attachment), tokens=getattr(raw, 'usage', {}), latency_ms=round((time.monotonic() - started) * 1000))
    if g.user:
        users.record_activity(g.user['id'], 'lesson')
    return jsonify({'response': response, 'messages': _strip_images(messages), 'done': 'Topic complete.' in response})

@app.route('/api/gamify')
def api_gamify():
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    return jsonify(users.get_gamify(g.user['id']))

@app.route('/api/reply', methods=['POST'])
def api_reply():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json
    messages = data.get('messages', [])
    user_input = data.get('input', '')
    user_input = user_input[:2000]
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    attachment = data.get('attachment')
    vision_model = data.get('vision_model', '')
    wrong_streak = int(data.get('wrong_streak', 0))
    tavily_key = _key(data, 'tavily_key', 'TAVILY_KEY')
    topic = data.get('topic', '')
    _, collection = _curriculum_params(data)
    if _mentions_papers(user_input):
        started = time.monotonic()
        intent = _paper_intent(user_input, topic, model=model, host=host, provider=provider, api_key=api_key)
        if intent['wants_papers']:
            reply, papers = _papers_for_chat(intent['topic'], tavily_key)
            messages.append({'role': 'user', 'content': user_input})
            messages.append({'role': 'assistant', 'content': reply})
            log_turn('papers', session_id=data.get('session_id', ''), user_id=g.user['id'], mode='learn', topic=intent['topic'] or topic, provider=provider, model=model, input=user_input, output=reply, correct=None, wrong_streak=wrong_streak, attachment=attachment_meta(attachment), tokens={}, latency_ms=round((time.monotonic() - started) * 1000))
            return jsonify({'response': reply, 'messages': _strip_images(messages), 'done': False, 'correct': None, 'papers': papers})
    detected_question = data.get('detected_question', '')
    user_msg = {'role': 'user', 'content': user_input + hint_note(wrong_streak)}
    if detected_question:
        user_msg['content'] += f"\n\n[Question from the student's image: {detected_question}]"
    if attachment:
        _apply_attachment(user_msg, attachment, provider, model, host, api_key, vision_model)
    messages.append(user_msg)
    started = time.monotonic()
    llm = dict(model=model, host=host, provider=provider, api_key=api_key)
    last_question = next((m['content'] for m in reversed(messages[:-1]) if m['role'] == 'assistant'), '')

    def _tutor_call():
        return _tutor_reply(chat(messages, **llm))

    def _eval_call():
        if not user_input.strip():
            print('EVAL: empty input, skipping')
            return None
        try:
            r = chat([{'role': 'user', 'content': f"A student was asked this question:\n{last_question}\n\nThe student replied: {user_input}\n\nIs the student's reply an attempt to answer the question (a value, calculation, or factual response)?\n- If YES and the answer is correct: reply CORRECT\n- If YES and the answer is wrong: reply WRONG\n- If NO (greeting, 'idk', frustration, off-topic, asking a question): reply SKIP\nReply with exactly one word."}], **llm)
            upper = str(r).upper()
            if 'SKIP' in upper:
                return None
            if 'WRONG' in upper:
                return False
            if 'CORRECT' in upper:
                return True
        except Exception as e:
            print(f'EVAL ERR: {e}')
        return None
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        tf = pool.submit(_tutor_call)
        ef = pool.submit(_eval_call)
        response, tutor_signal = tf.result()
        eval_signal = ef.result()
    correct = tutor_signal if tutor_signal is not None else eval_signal
    if correct is None and _AFFIRM_RE.match(response.strip()):
        correct = True
    messages.append({'role': 'assistant', 'content': response})
    log_turn('reply', session_id=data.get('session_id', ''), user_id=g.user['id'], mode='learn', topic=topic, provider=provider, model=model, input=user_input, output=response, correct=correct, wrong_streak=wrong_streak, attachment=attachment_meta(attachment), tokens={}, latency_ms=round((time.monotonic() - started) * 1000))
    cid = (data.get('curriculum_id') or '').strip()
    if correct is not None and g.user and cid and topic:
        users.record_progress(g.user['id'], cid, topic, correct)
        users.record_activity(g.user['id'], 'answer', correct=correct)
    return jsonify({'response': response, 'messages': _strip_images(messages), 'done': 'Topic complete.' in response, 'correct': correct})
_SESSION_ID_MAX = 64
_REPLAYABLE = ('start', 'reply', 'papers')

def _rebuild_messages(turns: list[dict]) -> list[dict]:
    messages = [{'role': 'system', 'content': TUTOR_SYSTEM_PROMPT}]
    for turn in turns:
        if turn.get('event') not in _REPLAYABLE:
            continue
        text = (turn.get('input') or '').strip()
        if turn.get('event') == 'start':
            text = f'Teach me about: {text}'
        meta = turn.get('attachment')
        if meta:
            text += f"\n\n[The student attached a {meta.get('type', 'file')}: {meta.get('name', '')}. It is not available in this resumed session.]"
        messages.append({'role': 'user', 'content': text})
        messages.append({'role': 'assistant', 'content': turn.get('output') or ''})
    return messages

@app.route('/api/sessions')
def api_sessions():
    return jsonify({'sessions': list_sessions(g.user['id'])})

@app.route('/api/sessions/<session_id>')
def api_session(session_id: str):
    if len(session_id) > _SESSION_ID_MAX:
        return (jsonify({'error': 'bad session id'}), 400)
    turns = session_turns(session_id, g.user['id'])
    if not turns:
        return (jsonify({'error': 'no such session'}), 404)
    lesson = [t for t in turns if t.get('event') in _REPLAYABLE]
    shown = lesson or turns
    return jsonify({'session_id': session_id, 'mode': turns[0].get('mode', 'learn'), 'topic': turns[0].get('topic', ''), 'turns': [{'event': t.get('event', ''), 'input': t.get('input', ''), 'output': t.get('output', ''), 'correct': t.get('correct')} for t in shown], 'messages': _rebuild_messages(lesson), 'resumable': bool(lesson)})

@app.route('/api/check', methods=['POST'])
def api_check():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json
    question = data.get('question', '')
    answer = data.get('answer', '')
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    attachment = data.get('attachment')
    vision_model = data.get('vision_model', '')
    system_prompt, collection = _curriculum_params(data)
    context_chunks = retrieve(question, collection_name=collection)
    context = '\n\n'.join(context_chunks) if context_chunks else ''
    context_prefix = f'Reference material:\n{context}\n\n' if context else ''
    if attachment and attachment.get('type') == 'image' and (provider == 'ollama'):
        messages = [{'role': 'system', 'content': system_prompt}, {'role': 'user', 'content': build_check_prompt(question, answer, context_prefix=context_prefix, note='The student also attached an image for context. '), 'images': [attachment['b64']]}]
        started = time.monotonic()
        raw = chat(messages, model=model, host=host, provider=provider, api_key=api_key)
        response, correct = _tutor_reply(raw)
        log_turn('check', session_id=data.get('session_id', ''), user_id=g.user['id'], mode='check', topic=question, provider=provider, model=model, input=answer, output=response, correct=correct, attachment=attachment_meta(attachment), tokens=getattr(raw, 'usage', {}), latency_ms=round((time.monotonic() - started) * 1000))
        return jsonify({'response': response, 'correct': correct})
    pdf_context = f"Attached document:\n{attachment['text']}\n\n" if attachment and attachment.get('type') == 'pdf' else ''
    image_context = ''
    if attachment and attachment.get('type') == 'image':
        image_context = _transcription_note(attachment, provider, model, host, api_key, vision_model).strip() + '\n\n'
    prompt = build_check_prompt(question, answer, context_prefix=context_prefix)
    started = time.monotonic()
    raw = generate(prompt, model=model, host=host, provider=provider, api_key=api_key)
    response, correct = _tutor_reply(raw)
    log_turn('check', session_id=data.get('session_id', ''), user_id=g.user['id'], mode='check', topic=question, provider=provider, model=model, input=answer, output=response, correct=correct, attachment=attachment_meta(attachment), tokens=getattr(raw, 'usage', {}), latency_ms=round((time.monotonic() - started) * 1000))
    return jsonify({'response': response, 'correct': correct})

@app.route('/api/check_working', methods=['POST'])
def api_check_working():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json or {}
    b64 = data.get('b64', '')
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    vision_model = data.get('vision_model', '')
    if not b64:
        return (jsonify({'error': 'no image provided'}), 400)
    img_bytes = base64.b64decode(b64)
    with tempfile.NamedTemporaryFile(suffix='.jpg', delete=False) as f:
        f.write(img_bytes)
        tmp_path = f.name
    try:
        started = time.monotonic()
        result = check_working(tmp_path, provider=provider, model=model, host=host, api_key=api_key, vision_model=vision_model)
        latency = round((time.monotonic() - started) * 1000)
    except Exception as e:
        return (jsonify({'error': str(e)}), 500)
    finally:
        Path(tmp_path).unlink(missing_ok=True)
    log_turn('check_working', session_id=data.get('session_id', ''), user_id=g.user['id'], provider=provider, model=model, steps=len(result['steps']), error_index=result['error_index'], latency_ms=latency)
    if g.user and result.get('error_index') is not None:
        cid = (data.get('curriculum_id') or '').strip()
        topic = (data.get('topic') or '').strip()
        if cid and topic:
            category = core.misconceptions.normalise(result.get('category'))
            users.record_misconception(g.user['id'], cid, topic, category)
    if g.user:
        users.record_activity(g.user['id'], 'check_working')
    return jsonify(result)

@app.route('/api/detect_question', methods=['POST'])
def api_detect_question():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json or {}
    b64 = data.get('b64', '')
    if not b64:
        return (jsonify({'error': 'no image'}), 400)
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    vision_model = data.get('vision_model', '')
    from core.vision.image_handler import detect_question
    try:
        result = detect_question(b64, name='image.jpg', provider=provider, model=model, host=host, api_key=api_key, vision_model=vision_model)
    except Exception as e:
        return (jsonify({'error': str(e)}), 500)
    return jsonify(result)

def _apply_attachment(msg: dict, attachment: dict, provider: str, model: str='', host: str='', api_key: str='', vision_model: str=''):
    if attachment.get('type') == 'image':
        name = attachment.get('name', 'image')
        if provider == 'ollama':
            msg['images'] = [attachment['b64']]
            msg['content'] += f'\n\n(Student attached image: {name})'
        else:
            msg['content'] += _transcription_note(attachment, provider, model, host, api_key, vision_model)
    elif attachment.get('type') == 'pdf':
        msg['content'] += f"\n\nAttached document ({attachment.get('name', 'file')}):\n{attachment['text']}"

def _transcription_note(attachment: dict, provider: str, model: str, host: str, api_key: str, vision_model: str) -> str:
    name = attachment.get('name', 'image')
    try:
        text = transcribe(attachment['b64'], name=name, provider=provider, model=model, host=host, api_key=api_key, vision_model=vision_model)
    except Exception as e:
        app.logger.warning('Image transcription failed for %s: %s', name, e)
        return f'\n\n(Student attached image: {name} — could not be read: {e})'
    return f'\n\nThe student attached an image ({name}). Transcription:\n{text}'

def _strip_images(messages: list) -> list:
    return [{'role': m['role'], 'content': m['content']} for m in messages]

def _ai_evaluate(student_input: str, tutor_reply: str, **llm) -> bool | None:
    msgs = [{'role': 'user', 'content': f"A tutor responded to a student.\n\nStudent said: {student_input}\n\nTutor replied: {tutor_reply}\n\nBased only on the tutor's reply, was the student's answer correct? Reply with exactly one word: CORRECT or WRONG."}]
    try:
        r = chat(msgs, **llm)
        upper = str(r).upper()
        if 'CORRECT' in upper:
            return True
        if 'WRONG' in upper:
            return False
    except Exception:
        pass
    return None
_NUMERIC_RE = re.compile('\\d')
_AFFIRM_RE = re.compile("^(right|correct|exactly|yes[,!.\\s]|great|perfect|spot.?on|well done|that\\'?s?\\s+(right|correct)|good\\s+(job|work|answer|one))", re.IGNORECASE)
_FILLER_RE = re.compile("\\b(is\\s+(it|this(\\s+the\\s+answer)?|that\\s+(right|correct))|i\\s+think\\s+(it\\'?s?|the\\s+answer\\s+is)|the\\s+answer\\s+is|would\\s+it\\s+be|so\\s+it\\'?s?)\\b.*$", re.IGNORECASE)

def _extract_math(text: str) -> str:
    cleaned = _FILLER_RE.sub('', text).strip().strip(',').strip()
    return cleaned if cleaned else text.strip()
_NUMBER_TOKEN_RE = re.compile('[-+]?\\d+(?:[./]\\d+)?')

def _parse_number(s: str):
    import fractions
    s = s.strip().rstrip('%').strip()
    try:
        return fractions.Fraction(s)
    except Exception:
        return None

def _cross_check(student_input: str, tutor_reply: str) -> bool | None:
    student_val = _parse_number(_extract_math(student_input))
    if student_val is None:
        return None
    nums = _NUMBER_TOKEN_RE.findall(tutor_reply)
    if not nums:
        return None
    tutor_val = _parse_number(nums[-1])
    if tutor_val is None:
        return None
    return None if student_val == tutor_val else False

def _verify_math(student_input: str, messages: list, **llm) -> bool | None:
    question = next((m['content'] for m in reversed(messages) if m['role'] == 'assistant'), '')
    prompt = f"A student answered this maths question:\n{question}\n\nTheir answer: {student_input}\n\nSolve it yourself from scratch. Ignore what anyone else said. Is the student's answer correct? Reply with exactly one word: VERIFIED or WRONG."
    try:
        r = chat([{'role': 'user', 'content': prompt}], **llm)
        upper = str(r).upper()
        if 'VERIFIED' in upper:
            return True
        if 'WRONG' in upper:
            return False
    except Exception:
        pass
    return None
PAPER_DOMAINS = ['physicsandmathstutor.com', 'savemyexams.co.uk', 'aqa.org.uk', 'cambridgeinternational.org', 'ocr.org.uk', 'qualifications.pearson.com', 'collegeboard.org', 'pastpapers.co', 'revisely.co.uk', 'znotes.org']
PAPER_GEN_PROMPT = "You are an exam question writer.\nRead the content below and generate {n} multiple-choice questions that test understanding of key concepts.\n\nUse this EXACT format — no deviations:\n\nQ1. Question text here\nA) First option\nB) Second option\nC) Third option\nD) Fourth option\n[ANSWER: B]\n\nRules:\n- Number questions Q1, Q2, Q3 etc.\n- Always exactly 4 options: A B C D\n- Always end each question with [ANSWER: X] on its own line\n- One correct answer, three plausible distractors\n- Blank line between questions\n- Every question must test a mathematical concept, fact, or procedure — not document structure.\n- NEVER ask about worksheet titles, document names, topic headings, or how content is organised.\n- NEVER ask about the exam board, the publisher, or the format of the notes.\n- BAD example (forbidden): 'Which topic includes a worksheet titled X?' — this tests nothing mathematical.\n- GOOD example: 'What is the gradient of the line y = 3x + 2?' — this tests mathematical knowledge.\n\nContent:\n{content}"
PRACTICE_FROM_ERROR_PROMPT = 'A student made a mistake in their maths working.\n\nThe line they got wrong: {step}\nWhat went wrong: {explanation}\n\nWrite {n} multiple-choice questions that give focused practice on exactly this kind of mistake, at the same topic and difficulty, so the student learns to avoid it.\n\nUse this EXACT format — no deviations:\n\nQ1. Question text here\nA) First option\nB) Second option\nC) Third option\nD) Fourth option\n[ANSWER: B]\n\nRules:\n- Number questions Q1, Q2, Q3 etc.\n- Always exactly 4 options: A B C D\n- Always end each question with [ANSWER: X] on its own line\n- One correct answer, three plausible distractors\n- Every question must test the mathematical skill the student got wrong.\n- Blank line between questions.\n'

@app.route('/api/practice_from_error', methods=['POST'])
def api_practice_from_error():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json or {}
    step = (data.get('step') or '').strip()
    explanation = (data.get('explanation') or '').strip()
    n = int(data.get('n_questions', 3))
    provider = data.get('provider', 'ollama')
    model = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    if not explanation:
        return (jsonify({'error': 'no error to practise'}), 400)
    prompt = PRACTICE_FROM_ERROR_PROMPT.format(n=n, step=step or '(their working)', explanation=explanation)
    try:
        raw = generate(prompt, model=model, host=host, provider=provider, api_key=api_key)
    except Exception as e:
        return (jsonify({'error': str(e)}), 500)
    questions = _parse_questions(str(raw))
    if not questions:
        return (jsonify({'error': 'could not generate practice questions'}), 502)
    return jsonify({'questions': _format_questions(questions)})

def _allowed_paper_host(url: str) -> bool:
    host = (urlparse(url).hostname or '').lower()
    return any((host == d or host.endswith('.' + d) for d in PAPER_DOMAINS))

def _fetch_pdf(url: str) -> bytes | None:
    if not _allowed_paper_host(url):
        return None
    try:
        with requests.get(url, timeout=20, stream=True) as resp:
            resp.raise_for_status()
            ctype = resp.headers.get('Content-Type', '').lower()
            if 'pdf' not in ctype and (not url.lower().endswith('.pdf')):
                return None
            chunks, total = ([], 0)
            for chunk in resp.iter_content(64 * 1024):
                total += len(chunk)
                if total > PAPER_MAX_BYTES:
                    return None
                chunks.append(chunk)
    except Exception:
        return None
    data = b''.join(chunks)
    return data if data.startswith(b'%PDF') else None

def _cache_paper(url: str, title: str, data: bytes) -> str:
    paper_id = hashlib.sha256(url.encode()).hexdigest()[:16]
    PAPER_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    (PAPER_CACHE_DIR / f'{paper_id}.pdf').write_bytes(data)
    (PAPER_CACHE_DIR / f'{paper_id}.json').write_text(json.dumps({'url': url, 'title': title}))
    return paper_id

@app.route('/api/papers/download/<paper_id>')
def api_paper_download(paper_id: str):
    if not _PAPER_ID_RE.match(paper_id):
        return (jsonify({'error': 'Bad paper id'}), 400)
    pdf_path = PAPER_CACHE_DIR / f'{paper_id}.pdf'
    if not pdf_path.is_file():
        return (jsonify({'error': 'Paper not in cache — run the search again'}), 404)
    meta_path = PAPER_CACHE_DIR / f'{paper_id}.json'
    title = json.loads(meta_path.read_text()).get('title', '') if meta_path.is_file() else ''
    name = secure_filename(title) or paper_id
    if not name.lower().endswith('.pdf'):
        name += '.pdf'
    return send_file(pdf_path, as_attachment=True, download_name=name, mimetype='application/pdf')
VERIFY_PROMPT = 'You are checking the answer key of a multiple-choice exam paper.\nSolve every question below from scratch and decide the correct option yourself. Work through the reasoning briefly for each one — do not guess from how the options look.\nIf a question is unanswerable, contradictory, or no option is correct, answer X for it.\n\nAfter your working, end your reply with a final block in exactly this format and nothing after it:\nFINAL\nQ1: B\nQ2: D\n\nQuestions:\n{questions}'
_Q_BLOCK_RE = re.compile('\\n(?=Q\\d+\\.)')
_Q_OPTION_RE = re.compile('^([A-D])\\)\\s*(.+)$', re.MULTILINE)
_Q_ANSWER_RE = re.compile('\\[ANSWER:\\s*([A-D])\\]', re.IGNORECASE)
_VERDICT_RE = re.compile('^Q(\\d+)\\s*:\\s*([A-DX])\\s*$', re.MULTILINE | re.IGNORECASE)

def _parse_questions(raw: str) -> list[dict]:
    questions = []
    for block in _Q_BLOCK_RE.split(raw or ''):
        if not block.strip():
            continue
        options = _Q_OPTION_RE.findall(block)
        answer = _Q_ANSWER_RE.search(block)
        if len(options) < 2 or not answer:
            continue
        text = _Q_OPTION_RE.split(block)[0]
        text = re.sub('^Q\\d+\\.\\s*', '', text).strip()
        if text:
            questions.append({'text': text, 'options': options, 'answer': answer.group(1).upper()})
    return questions

def _format_questions(questions: list[dict]) -> str:
    out = []
    for i, q in enumerate(questions, 1):
        lines = [f"Q{i}. {q['text']}"]
        lines += [f'{letter}) {body}' for letter, body in q['options']]
        lines.append(f"[ANSWER: {q['answer']}]")
        out.append('\n'.join(lines))
    return '\n\n'.join(out)

def _verify_questions(questions: list[dict], **llm) -> tuple[list[dict], int]:
    if not questions:
        return (questions, 0)
    asked = '\n\n'.join(('\n'.join([f"Q{i}. {q['text']}"] + [f'{letter}) {body}' for letter, body in q['options']]) for i, q in enumerate(questions, 1)))
    try:
        reply = chat([{'role': 'user', 'content': VERIFY_PROMPT.format(questions=asked)}], **llm)
    except Exception:
        return (questions, 0)
    tail = reply.rsplit('FINAL', 1)[-1]
    verdicts = {int(n): letter.upper() for n, letter in _VERDICT_RE.findall(tail)}
    kept = []
    for i, q in enumerate(questions, 1):
        verdict = verdicts.get(i)
        if verdict is None or verdict == q['answer']:
            kept.append(q)
    return (kept, len(questions) - len(kept))

def _rag_questions(topic: str, data: dict, n_questions: int):
    _, collection = _curriculum_params(data)
    chunks = retrieve(topic, n_results=10, collection_name=collection, with_sources=True)
    if not chunks:
        return ([], [], 0, (jsonify({'error': 'No curriculum content found for that topic. Add a Tavily key or ingest reference material first.'}), 404))
    combined = '\n\n---\n\n'.join((c['text'] for c in chunks))[:7000]
    seen_titles: set[str] = set()
    sources = []
    for c in chunks:
        t = c['source']
        if t and t not in seen_titles:
            seen_titles.add(t)
            sources.append({'title': t, 'url': _pmt_url(t)})
    llm_kwargs = {'model': data.get('model', 'gemma4'), 'host': data.get('host', 'http://localhost:11434'), 'provider': data.get('provider', 'ollama'), 'api_key': _key(data, 'api_key', 'API_KEY')}
    prompt = PAPER_GEN_PROMPT.format(n=n_questions + VERIFY_OVERSHOOT, content=combined)
    try:
        raw = chat([{'role': 'user', 'content': prompt}], **llm_kwargs)
    except Exception as e:
        return ([], sources, 0, (jsonify({'error': f'Question generation failed: {e}'}), 500))
    questions = _parse_questions(raw)
    verified, dropped = _verify_questions(questions, **llm_kwargs)
    final = (verified if verified else questions)[:n_questions]
    if not final:
        return ([], sources, dropped, (jsonify({'error': 'Could not generate questions for that topic. Try a more specific topic.'}), 502))
    return (final, sources, dropped, None)

def _attempt_questions(raw_questions: list) -> list[dict]:
    out = []
    for i, q in enumerate(raw_questions, 1):
        options = []
        for opt in q.get('options') or []:
            if isinstance(opt, dict):
                options.append({'letter': (opt.get('letter') or '').upper(), 'text': opt.get('text', '')})
            else:
                letter, body = opt
                options.append({'letter': letter.upper(), 'text': body})
        out.append({'id': str(q.get('id') or f'q{i}'), 'text': q.get('text', ''), 'options': options, 'answer': (q.get('answer') or '').upper(), 'working': q.get('working', '')})
    return out

def _public_questions(questions: list[dict]) -> list[dict]:
    return [{'id': q['id'], 'text': q['text'], 'options': q['options']} for q in questions]

def _attempt_or_404(data: dict):
    attempt = users.get_attempt((data.get('attempt_id') or '').strip(), g.user['id'])
    if attempt is None:
        return (None, (jsonify({'error': 'Unknown attempt'}), 404))
    return (attempt, None)

@app.route('/api/quiz/start', methods=['POST'])
def api_quiz_start():
    data = request.json or {}
    cid = (data.get('curriculum_id') or '').strip()
    topic = (data.get('topic') or '').strip() or cid
    n = max(1, min(int(data.get('n_questions', 5)), 20))
    if not cid:
        return (jsonify({'error': 'No curriculum given'}), 400)
    source = 'bank'
    questions = _attempt_questions(quizbank.pick(cid, n))
    if not questions:
        err = _check_llm_rate()
        if err:
            return err
        generated, _sources, _dropped, gen_err = _rag_questions(topic, data, n)
        if gen_err:
            return gen_err
        questions = _attempt_questions(generated)
        source = 'generated'
    if not questions:
        return (jsonify({'error': 'No quiz questions available for this class yet.'}), 404)
    attempt_id = users.create_attempt(g.user['id'], cid, topic, questions)
    return jsonify({'attempt_id': attempt_id, 'source': source, 'questions': _public_questions(questions)})

@app.route('/api/quiz/answer', methods=['POST'])
def api_quiz_answer():
    data = request.json or {}
    attempt, err = _attempt_or_404(data)
    if err:
        return err
    qid = (data.get('qid') or '').strip()
    choice = (data.get('choice') or '').strip().upper()
    question = next((q for q in attempt['questions'] if q['id'] == qid), None)
    if question is None:
        return (jsonify({'error': 'Unknown question'}), 404)
    correct = choice == question['answer']
    if qid not in attempt['marked']:
        users.record_progress(g.user['id'], attempt['curriculum'], attempt['topic'], correct)
        users.record_activity(g.user['id'], 'answer', correct=correct)
    users.record_answer(attempt['id'], g.user['id'], qid, choice, correct)
    return jsonify({'correct': correct, 'correct_answer': question['answer'], 'working': question.get('working', '')})

@app.route('/api/quiz/tab-switch', methods=['POST'])
def api_quiz_tab_switch():
    data = request.json or {}
    attempt, err = _attempt_or_404(data)
    if err:
        return err
    return jsonify({'ok': True, 'tab_switches': users.bump_tab_switch(attempt['id'], g.user['id'])})

@app.route('/api/quiz/submit', methods=['POST'])
def api_quiz_submit():
    data = request.json or {}
    attempt, err = _attempt_or_404(data)
    if err:
        return err
    sent = {str(k): str(v).upper() for k, v in (data.get('answers') or {}).items()}
    resubmit = attempt['submitted_at'] is not None
    score = 0
    review = []
    for question in attempt['questions']:
        qid = question['id']
        chosen = sent.get(qid) or attempt['answers'].get(qid, '')
        ok = bool(chosen) and chosen == question['answer']
        if ok:
            score += 1
        if not resubmit and qid not in attempt['marked']:
            users.record_progress(g.user['id'], attempt['curriculum'], attempt['topic'], ok)
            users.record_activity(g.user['id'], 'answer', correct=ok)
            users.record_answer(attempt['id'], g.user['id'], qid, chosen, ok)
        review.append({'id': qid, 'text': question['text'], 'options': question['options'], 'chosen': chosen, 'correct_answer': question['answer'], 'correct': ok, 'working': question.get('working', '')})
    if not resubmit:
        users.finish_attempt(attempt['id'], g.user['id'], score)
    return jsonify({'score': score, 'total': len(attempt['questions']), 'tab_switches': attempt['tab_switches'], 'resubmitted': resubmit, 'review': review})

@app.route('/api/papers', methods=['POST'])
def api_papers():
    err = _check_llm_rate()
    if err:
        return err
    data = request.json
    topic = data.get('topic', '')
    tavily_key = _key(data, 'tavily_key', 'TAVILY_KEY')
    n_questions = int(data.get('n_questions', 5))
    provider = data.get('provider', 'ollama')
    model_name = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    if not topic:
        return (jsonify({'error': 'No topic provided'}), 400)
    if not tavily_key:
        final, sources, dropped, gen_err = _rag_questions(topic, data, n_questions)
        if gen_err:
            return gen_err
        return jsonify({'questions': _format_questions(final), 'sources': sources, 'papers': [], 'grounded': True, 'dropped': dropped, 'attempt_id': _open_paper_attempt(data, topic, final)})
    try:
        results = _search_papers(topic, tavily_key)
    except Exception as e:
        return (jsonify({'error': f'Search failed: {e}'}), 500)
    if not results:
        return (jsonify({'error': 'No past paper content found for that topic. Try a more specific query.'}), 404)
    papers, paper_texts = _collect_papers(results)
    if paper_texts:
        combined = '\n\n---\n\n'.join(paper_texts)[:7000]
    else:
        content_parts = [r['content'] for r in results[:5] if r.get('content')]
        combined = '\n\n---\n\n'.join(content_parts)[:7000]
    sources = [{'title': r.get('title', ''), 'url': r.get('url', '')} for r in results[:5]]
    prompt = PAPER_GEN_PROMPT.format(n=n_questions + VERIFY_OVERSHOOT, content=combined)
    try:
        raw = chat([{'role': 'user', 'content': prompt}], model=model_name, host=host, provider=provider, api_key=api_key)
    except Exception as e:
        return (jsonify({'error': f'Question generation failed: {e}'}), 500)
    llm_kwargs = {'model': model_name, 'host': host, 'provider': provider, 'api_key': api_key}
    questions = _parse_questions(raw)
    verified, dropped = _verify_questions(questions, **llm_kwargs)
    verified = verified[:n_questions]
    if not verified:
        return (jsonify({'error': 'Generated questions all failed the answer check. Try a more specific topic.'}), 502)
    return jsonify({'questions': _format_questions(verified), 'sources': sources, 'papers': papers, 'grounded': bool(paper_texts), 'dropped': dropped, 'attempt_id': _open_paper_attempt(data, topic, verified)})

def _open_paper_attempt(data: dict, topic: str, questions: list[dict]) -> str:
    cid = (data.get('curriculum_id') or '').strip()
    if not cid or not questions:
        return ''
    return users.create_attempt(g.user['id'], cid, topic or cid, _attempt_questions(questions))

@app.route('/api/papers/hint', methods=['POST'])
def api_papers_hint():
    data = request.json
    question = data.get('question', '')
    options = data.get('options', [])
    chosen = data.get('chosen', '')
    correct = data.get('correct', '')
    wrong_streak = int(data.get('wrong_streak', 1))
    provider = data.get('provider', 'ollama')
    model_name = data.get('model', 'gemma4')
    host = data.get('host', 'http://localhost:11434')
    api_key = _key(data, 'api_key', 'API_KEY')
    if wrong_streak == 1 and g.user:
        cid = (data.get('curriculum_id') or '').strip()
        topic = (data.get('topic') or '').strip()
        if cid and topic:
            users.record_progress(g.user['id'], cid, topic, correct=False)
            users.record_activity(g.user['id'], 'answer', correct=False)
    opts_text = '\n'.join((f"{o['letter']}) {o['text']}" for o in options))
    if wrong_streak == 1:
        instruction = 'Give a subtle hint that nudges the student toward the right thinking. Stay indirect — do not name the concept yet.'
    elif wrong_streak == 2:
        instruction = 'Name the key concept the student needs but do not give away the answer.'
    else:
        instruction = 'Directly explain why the chosen option is wrong and explain the correct concept clearly. Then ask the student to confirm they understand.'
    messages = [{'role': 'system', 'content': TUTOR_SYSTEM_PROMPT}, {'role': 'user', 'content': f'Past paper question:\n{question}\n\n{opts_text}\n\nThe student chose: {chosen}\nThe correct answer is: {correct}\n\n{instruction} Write in plain, simple English. Do not end with [CORRECT] or [WRONG].'}]
    response = chat(messages, model=model_name, host=host, provider=provider, api_key=api_key)
    response, _ = _tutor_reply(response)
    return jsonify({'response': response})

def _pmt_url(source_title: str) -> str:
    import re
    parts = source_title.split(' - ', 1)
    if len(parts) == 2:
        clean = re.sub('[^\\w\\s]', '', parts[1].strip().lower())
        slug = re.sub('\\s+', '-', clean)
        return f'https://www.physicsandmathstutor.com/maths-revision/gcse-maths/{slug}/'
    return 'https://www.physicsandmathstutor.com/maths-revision/gcse-maths/'

def _search_papers(topic: str, tavily_key: str) -> list[dict]:
    resp = requests.post('https://api.tavily.com/search', json={'api_key': tavily_key, 'query': f'{topic} past paper exam questions SAT A-level filetype:pdf', 'search_depth': 'advanced', 'include_domains': PAPER_DOMAINS, 'include_raw_content': True, 'max_results': 10}, timeout=20)
    resp.raise_for_status()
    return resp.json().get('results', [])

def _collect_papers(results: list[dict]) -> tuple[list[dict], list[str]]:
    papers, paper_texts = ([], [])
    for r in results:
        if len(papers) >= PAPER_MAX_FETCH:
            break
        url = r.get('url', '')
        if not url.lower().endswith('.pdf'):
            continue
        data_bytes = _fetch_pdf(url)
        if not data_bytes:
            continue
        title = r.get('title', '') or url.rsplit('/', 1)[-1]
        try:
            has_text, text = _question_pages(data_bytes)
            text = text.strip()
        except Exception:
            has_text, text = (False, '')
        paper_id = _cache_paper(url, title, data_bytes)
        papers.append({'id': paper_id, 'title': title, 'url': url, 'download': f'/api/papers/download/{paper_id}', 'scanned': not has_text, 'no_questions': has_text and (not text)})
        if text:
            paper_texts.append(text)
    return (papers, paper_texts)
_PAPER_NOUN = '(?:past[\\s-]?papers?|exam papers?|pdfs?|mark schemes?|real (?:exam )?questions?|papers)'
_FETCH_VERB = '(?:get|give|fetch|find|download|send|show|grab|take|provide|need|want|have)'
_NEGATION = "(?:don'?t|do not|not|no need|nope|never|rather not|without)"
_ABOUT_RE = re.compile('^\\s*(?:what|why|how|explain|describe|tell me about)\\b', re.IGNORECASE)
_WANTS_PAPERS_RE = re.compile(f'\\b{_FETCH_VERB}\\b[^.?!]{{0,40}}?\\b{_PAPER_NOUN}\\b', re.IGNORECASE)
_NOT_PAPERS_RE = re.compile(f'\\b{_NEGATION}\\b[^.?!]{{0,40}}?\\b{_PAPER_NOUN}\\b', re.IGNORECASE)

def _wants_papers(text: str) -> bool:
    if not text or _NOT_PAPERS_RE.search(text):
        return False
    if _WANTS_PAPERS_RE.search(text):
        return True
    if _ABOUT_RE.search(text):
        return False
    return len(text.split()) <= 8 and bool(re.search(_PAPER_NOUN, text, re.I))
_TOPIC_IN_REQUEST_RE = re.compile(f"{_PAPER_NOUN}\\s+(?:on|for|about|in)\\s+([a-z0-9][a-z0-9 '\\-]{{2,50}})", re.IGNORECASE)
_TOPIC_TRAIL_RE = re.compile('\\b(?:please|pls|thanks|thank you|now|asap|ok)\\b\\s*$', re.IGNORECASE)

def _stated_topic(user_input: str) -> str:
    m = _TOPIC_IN_REQUEST_RE.search(user_input)
    return _TOPIC_TRAIL_RE.sub('', m.group(1)).strip() if m else ''
_MENTIONS_PAPERS_RE = re.compile(f'\\b(?:{_PAPER_NOUN}|paper)\\b', re.IGNORECASE)

def _mentions_papers(text: str) -> bool:
    return bool(text) and bool(_MENTIONS_PAPERS_RE.search(text))
PAPER_INTENT_PROMPT = 'You decide whether a student is asking to be given past exam paper files, and what subject those papers should be about.\n\nReturn ONLY a JSON object. No prose, no code fences:\n{{"wants_papers": true or false, "topic": "subject" or null}}\n\nwants_papers is true when the student wants actual paper files handed to them, including when it is phrased as a complaint or a rhetorical question ("can you not just fetch the whole pdf", "why won\'t you give me the paper").\nwants_papers is false when they ask what a past paper is, discuss a paper they already uploaded, or are simply answering a lesson question.\n\ntopic is the exam subject to search for, as a short phrase a search engine would accept: "SAT", "trigonometry", "AS level Computer Science". Take it from the student\'s message when they name one. Otherwise fall back to the lesson topic below, shortened to just the subject — the lesson topic may be a full sentence, so strip it down to the subject. Use null only if neither gives a subject.\n\nLesson topic: {lesson_topic}\nStudent message: {message}'
_JSON_OBJ_RE = re.compile('\\{.*\\}', re.DOTALL)

def _paper_intent(user_input: str, lesson_topic: str, **llm) -> dict:
    fallback = {'wants_papers': _wants_papers(user_input), 'topic': _stated_topic(user_input) or lesson_topic.strip()}
    try:
        raw = chat([{'role': 'user', 'content': PAPER_INTENT_PROMPT.format(lesson_topic=lesson_topic.strip() or '(none)', message=user_input)}], **llm)
    except Exception:
        return fallback
    m = _JSON_OBJ_RE.search(raw or '')
    if not m:
        return fallback
    try:
        data = json.loads(m.group(0))
    except (ValueError, TypeError):
        return fallback
    if not isinstance(data.get('wants_papers'), bool):
        return fallback
    topic = data.get('topic')
    topic = topic.strip() if isinstance(topic, str) else ''
    return {'wants_papers': data['wants_papers'], 'topic': topic}
_TOPIC_PHRASE_MAX_WORDS = 6
_TOPIC_PHRASE_MAX_CHARS = 40

def _topic_phrase(topic: str) -> str:
    t = ' '.join((topic or '').split())
    if not t or len(t) > _TOPIC_PHRASE_MAX_CHARS or len(t.split()) > _TOPIC_PHRASE_MAX_WORDS:
        return ''
    return f' on {t}'

def _papers_for_chat(topic: str, tavily_key: str) -> tuple[str, list]:
    topic = (topic or '').strip()
    if topic and (not _looks_like_topic(topic)):
        return (f'''I'd be searching for "{topic}", which doesn't look like a subject. Which topic should I find papers on?''', [])
    if not topic:
        return ("Which topic should I look for? Tell me the subject and I'll find papers on it.", [])
    if not tavily_key:
        return ("I can't search for papers — no Tavily API key is set. Add one in the settings bar or in the .env file.", [])
    try:
        results = _search_papers(topic, tavily_key)
    except Exception as e:
        return (f'The search for papers failed: {e}', [])
    papers, _ = _collect_papers(results)
    if not papers:
        return (f"I couldn't find a downloadable paper{_topic_phrase(topic)}. Try naming the exam board or a narrower subject.", [])
    plural = 'paper' if len(papers) == 1 else 'papers'
    reply = f"Here {('is' if len(papers) == 1 else 'are')} {len(papers)} {plural}{_topic_phrase(topic)} — click to download."
    scans = sum((1 for p in papers if p['scanned']))
    if scans:
        reply += f" {scans} of them {('is a scan' if scans == 1 else 'are scans')}, so the text can't be read out of it — but it still downloads."
    reply += ' Your lesson is still here whenever you want to carry on.'
    return (reply, papers)

def _looks_like_topic(text: str) -> bool:
    t = text.strip()
    if len(t) < 3 or _GREETING_RE.match(t):
        return False
    words = re.findall("[a-z0-9']+", t.lower())
    return bool(words) and any((w not in _FILLER_WORDS for w in words))
import curricula as _cur

@app.route('/api/curricula')
def api_student_curricula():
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    cfg = _settings.load()
    result = []
    for cid in _cur.list_curricula():
        if cid == 'demo' and (not cfg['stub_testing']):
            continue
        try:
            m = _cur.load(cid)
            result.append({'id': cid, 'name': m.get('name', cid), 'level': m.get('level', ''), 'subjects': m.get('subjects', [])})
        except Exception:
            result.append({'id': cid, 'name': cid, 'level': '', 'subjects': []})
    return jsonify(result)

@app.route('/api/curricula/<cid>/questions')
def api_student_questions(cid):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    path = Path('curricula') / cid / 'questions.json'
    if not path.exists():
        return jsonify([])
    return jsonify(json.loads(path.read_text(encoding='utf-8')))

@app.route('/api/curricula/<cid>/questions', methods=['POST'])
def api_student_questions_add(cid):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    data = request.get_json(force=True) or {}
    text = (data.get('text') or '').strip()
    if not text:
        return (jsonify({'error': 'text required'}), 400)
    path = Path('curricula') / cid / 'questions.json'
    questions = json.loads(path.read_text(encoding='utf-8')) if path.exists() else []
    q = {'id': secrets.token_hex(8), 'text': text, 'answer': (data.get('answer') or '').strip(), 'marks': 1, 'image': data.get('image') or '', 'author': g.user['username']}
    questions.append(q)
    path.write_text(json.dumps(questions, indent=2, ensure_ascii=False), encoding='utf-8')
    return (jsonify(q), 201)

@app.route('/api/curricula/<cid>/info')
def api_student_curriculum_info(cid):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'Not found'}), 404)
    topics_path = Path('curricula') / cid / 'topics.json'
    topics = json.loads(topics_path.read_text(encoding='utf-8')) if topics_path.exists() else []
    return jsonify({'manifest': manifest, 'topics': topics})

@app.route('/api/progress/<cid>')
def api_get_progress(cid):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    return jsonify(users.get_progress(g.user['id'], cid))

@app.route('/api/teacher/curricula/<cid>/attempts')
def api_teacher_attempts(cid):
    err = _teacher_or_admin()
    if err:
        return err
    return jsonify(users.list_attempts(cid))

@app.route('/api/teacher/curricula/<cid>/progress')
def api_teacher_progress(cid):
    err = _teacher_or_admin()
    if err:
        return err
    return jsonify(users.get_all_progress(cid))

@app.route('/api/teacher/curricula/<cid>/misconceptions')
def api_teacher_misconceptions(cid):
    err = _teacher_or_admin()
    if err:
        return err
    return jsonify(users.get_misconceptions(cid))

@app.route('/api/curricula/<cid>/materials')
def api_student_materials(cid):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    ref_dir = Path('curricula') / cid / 'reference'
    if not ref_dir.exists():
        return jsonify([])
    files = [{'name': f.name, 'url': f'/api/curricula/{cid}/materials/{f.name}'} for f in sorted(ref_dir.iterdir()) if f.is_file()]
    return jsonify(files)

@app.route('/api/curricula/<cid>/materials/<filename>')
def api_student_material_file(cid, filename):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    ref_dir = Path('curricula') / cid / 'reference'
    return send_from_directory(ref_dir.resolve(), secure_filename(filename))

@app.route('/api/admin/curricula/<cid>/auto-research', methods=['POST'])
def api_admin_auto_research(cid):
    err = _admin_only()
    if err:
        return err
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'Curriculum not found'}), 404)
    tavily_key = os.environ.get('TAVILY_API_KEY', '').strip()
    if not tavily_key:
        return (jsonify({'error': 'TAVILY_API_KEY not set in .env'}), 400)
    api_key = os.environ.get('API_KEY', '').strip()
    name = manifest.get('name', cid)
    flavour = manifest.get('prompt_flavour', name)
    subjects = manifest.get('subjects', [])
    queries = [f'{name} curriculum overview objectives', f'{name} topics syllabus structure']
    for subj in subjects[:3]:
        queries.append(f'{name} {subj} key concepts')
    snippets = []
    for q in queries:
        try:
            res = requests.post('https://api.tavily.com/search', json={'api_key': tavily_key, 'query': q, 'search_depth': 'basic', 'max_results': 4}, timeout=15)
            if res.ok:
                for r in res.json().get('results', []):
                    content = (r.get('content') or '').strip()
                    if content:
                        snippets.append(f"[{r.get('title', '')}]\n{content}")
        except Exception:
            pass
    if not snippets:
        return (jsonify({'error': 'No search results returned'}), 502)
    combined = '\n\n---\n\n'.join(snippets)
    synthesis_prompt = f'You are writing reference notes for a tutor AI that teaches {flavour}. Using only the search results below, write clear structured notes covering: overview, main topics, key concepts, and learning objectives. Use plain English. Use markdown headings. Do not invent facts not in the sources.\n\nSearch results:\n{combined[:12000]}'
    try:
        notes = generate(synthesis_prompt, model='openai/gpt-oss-120b', provider='groq', api_key=api_key)
    except Exception as e:
        return (jsonify({'error': f'LLM synthesis failed: {e}'}), 502)
    ref_dir = Path('curricula') / cid / 'reference'
    ref_dir.mkdir(parents=True, exist_ok=True)
    (ref_dir / 'auto-research.txt').write_text(notes, encoding='utf-8')
    import subprocess
    subprocess.Popen(['python', '-m', 'rag.ingest', '--curriculum', cid], cwd=Path(__file__).parent)
    return jsonify({'ok': True, 'notes_length': len(notes), 'queries': queries, 'snippets_found': len(snippets)})
_TEACHER_ALLOWED_EXT = {'.pdf', '.txt', '.pptx'}
_TEACHER_MAX_BYTES = 32 * 1024 * 1024

def _teacher_or_admin():
    return _require_role('teacher', 'admin')

def _resolve_ref_dir(manifest: dict) -> Path:
    ref = _cur.reference_dir(manifest)
    p = Path(ref) if Path(ref).is_absolute() else Path(__file__).parent / ref
    p.mkdir(parents=True, exist_ok=True)
    return p

def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                parts.append(shape.text.strip())
    return '\n\n'.join(parts)

def _questions_path(cid: str) -> Path:
    return _cur.CURRICULA_DIR / cid / 'questions.json'

def _load_questions(cid: str) -> list:
    p = _questions_path(cid)
    return json.loads(p.read_text()) if p.exists() else []

def _save_questions(cid: str, questions: list) -> None:
    _questions_path(cid).write_text(json.dumps(questions, indent=2, ensure_ascii=False))

@app.route('/teacher')
def teacher_panel():
    return send_from_directory('static', 'teacher.html')

@app.route('/api/teacher/curricula')
def api_teacher_curricula():
    err = _teacher_or_admin()
    if err:
        return err
    result = []
    for cid in _cur.list_curricula():
        try:
            m = _cur.load(cid)
            result.append({'id': cid, 'name': m.get('name', cid)})
        except Exception:
            result.append({'id': cid, 'name': cid})
    return jsonify(result)

@app.route('/api/teacher/curricula/<cid>')
def api_teacher_get(cid):
    err = _teacher_or_admin()
    if err:
        return err
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    topics_path = _cur.CURRICULA_DIR / cid / 'topics.json'
    topics = json.loads(topics_path.read_text()) if topics_path.exists() else []
    ref_dir = _resolve_ref_dir(manifest)
    files = sorted(({'name': f.name, 'size': f.stat().st_size, 'ext': f.suffix.lower()} for f in ref_dir.iterdir() if f.is_file()))
    questions = _load_questions(cid)
    return jsonify({'manifest': {'id': manifest['id'], 'name': manifest['name']}, 'topics': topics, 'files': files, 'question_count': len(questions)})

@app.route('/api/teacher/curricula/<cid>/topics', methods=['PUT'])
def api_teacher_save_topics(cid):
    err = _teacher_or_admin()
    if err:
        return err
    try:
        _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    topics = request.json if isinstance(request.json, list) else []
    topics_path = _cur.CURRICULA_DIR / cid / 'topics.json'
    topics_path.write_text(json.dumps(topics, indent=2, ensure_ascii=False))
    return jsonify({'ok': True})

@app.route('/api/teacher/curricula/<cid>/upload', methods=['POST'])
def api_teacher_upload(cid):
    err = _teacher_or_admin()
    if err:
        return err
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    file = request.files.get('file')
    if not file:
        return (jsonify({'error': 'no file'}), 400)
    filename = secure_filename(file.filename or 'upload')
    ext = Path(filename).suffix.lower()
    if ext not in _TEACHER_ALLOWED_EXT:
        return (jsonify({'error': f"Allowed types: {', '.join(_TEACHER_ALLOWED_EXT)}"}), 400)
    data = file.read()
    if len(data) > _TEACHER_MAX_BYTES:
        return (jsonify({'error': 'File too large (max 32 MB)'}), 413)
    ref_dir = _resolve_ref_dir(manifest)
    if ext == '.pptx':
        text = _extract_pptx(data)
        if not text.strip():
            return (jsonify({'error': 'No readable text found in PowerPoint'}), 400)
        save_name = Path(filename).stem + '.txt'
        (ref_dir / save_name).write_text(text, encoding='utf-8')
    else:
        (ref_dir / filename).write_bytes(data)
        save_name = filename
    _start_ingest(cid, manifest)
    return jsonify({'ok': True, 'saved_as': save_name})

@app.route('/api/teacher/curricula/<cid>/files')
def api_teacher_files(cid):
    err = _teacher_or_admin()
    if err:
        return err
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    ref_dir = _resolve_ref_dir(manifest)
    files = sorted(({'name': f.name, 'size': f.stat().st_size} for f in ref_dir.iterdir() if f.is_file()))
    return jsonify(files)

@app.route('/api/teacher/curricula/<cid>/files/<filename>', methods=['DELETE'])
def api_teacher_delete_file(cid, filename):
    err = _teacher_or_admin()
    if err:
        return err
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    safe = secure_filename(filename)
    target = _resolve_ref_dir(manifest) / safe
    if not target.exists():
        return (jsonify({'error': 'file not found'}), 404)
    target.unlink()
    return jsonify({'ok': True})

@app.route('/api/teacher/curricula/<cid>/questions')
def api_teacher_questions(cid):
    err = _teacher_or_admin()
    if err:
        return err
    if not (_cur.CURRICULA_DIR / cid / 'manifest.json').exists():
        return (jsonify({'error': 'not found'}), 404)
    return jsonify(_load_questions(cid))

@app.route('/api/teacher/curricula/<cid>/questions', methods=['POST'])
def api_teacher_add_question(cid):
    err = _teacher_or_admin()
    if err:
        return err
    if not (_cur.CURRICULA_DIR / cid / 'manifest.json').exists():
        return (jsonify({'error': 'not found'}), 404)
    data = request.json or {}
    q_text = (data.get('question') or '').strip()
    answer = (data.get('answer') or '').strip()
    if not q_text or not answer:
        return (jsonify({'error': 'question and answer required'}), 400)
    questions = _load_questions(cid)
    new_q = {'id': secrets.token_hex(8), 'topic': (data.get('topic') or '').strip(), 'question': q_text, 'answer': answer, 'type': data.get('type', 'short_answer'), 'created_by': g.user['id'], 'created_at': int(time.time())}
    questions.append(new_q)
    _save_questions(cid, questions)
    return (jsonify(new_q), 201)

@app.route('/api/teacher/curricula/<cid>/questions/<qid>', methods=['DELETE'])
def api_teacher_delete_question(cid, qid):
    err = _teacher_or_admin()
    if err:
        return err
    questions = _load_questions(cid)
    questions = [q for q in questions if q.get('id') != qid]
    _save_questions(cid, questions)
    return jsonify({'ok': True})

def _start_ingest(cid: str, manifest: dict) -> bool:
    import threading
    from core.rag.ingest import build
    if _INGEST_STATUS.get(cid, {}).get('state') == 'running':
        return False
    _INGEST_STATUS[cid] = {'state': 'running', 'chunks': 0}

    def _run():
        try:
            n = build(text_folder=_cur.reference_dir(manifest), pdf_folder=_cur.papers_dir(manifest), collection_name=manifest['collection'])
            _INGEST_STATUS[cid] = {'state': 'idle', 'chunks': n}
        except Exception as exc:
            app.logger.error('Ingest failed for %s: %s', cid, exc)
            _INGEST_STATUS[cid] = {'state': 'error', 'chunks': 0}
    threading.Thread(target=_run, daemon=True).start()
    return True

@app.route('/api/teacher/curricula/<cid>/ingest', methods=['GET', 'POST'])
def api_teacher_ingest(cid):
    err = _teacher_or_admin()
    if err:
        return err
    if request.method == 'GET':
        return jsonify(_INGEST_STATUS.get(cid, {'state': 'idle', 'chunks': 0}))
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    started = _start_ingest(cid, manifest)
    return jsonify({'ok': True, 'already_running': not started})

def _require_role(*roles):
    if g.user is None:
        return (jsonify({'error': 'Not logged in'}), 401)
    if g.user.get('role') not in roles:
        return (jsonify({'error': 'Insufficient permissions'}), 403)
    return None

def _admin_only():
    return _require_role('admin')

@app.route('/admin')
def admin_panel():
    return send_from_directory('static', 'admin.html')

@app.route('/api/admin/curricula')
def api_admin_list():
    err = _admin_only()
    if err:
        return err
    result = []
    for cid in _cur.list_curricula():
        try:
            m = _cur.load(cid)
            result.append({'id': cid, 'name': m.get('name', cid)})
        except Exception:
            result.append({'id': cid, 'name': cid})
    return jsonify(result)

@app.route('/api/admin/curricula/<cid>', methods=['GET'])
def api_admin_get(cid):
    err = _admin_only()
    if err:
        return err
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    topics_path = _cur.CURRICULA_DIR / cid / 'topics.json'
    topics = json.loads(topics_path.read_text()) if topics_path.exists() else []
    return jsonify({'manifest': manifest, 'topics': topics})

@app.route('/api/admin/curricula/<cid>', methods=['PUT'])
def api_admin_save(cid):
    err = _admin_only()
    if err:
        return err
    data = request.json or {}
    manifest_path = _cur.CURRICULA_DIR / cid / 'manifest.json'
    if not manifest_path.exists():
        return (jsonify({'error': 'not found'}), 404)
    manifest = json.loads(manifest_path.read_text())
    for field in ('name', 'prompt_flavour', 'collection', 'reference_dir', 'papers_dir', 'level', 'subjects'):
        if field in data:
            manifest[field] = data[field]
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False))
    if 'topics' in data:
        topics_path = _cur.CURRICULA_DIR / cid / 'topics.json'
        topics_path.write_text(json.dumps(data['topics'], indent=2, ensure_ascii=False))
    return jsonify({'ok': True})

@app.route('/api/admin/curricula', methods=['POST'])
def api_admin_create():
    err = _admin_only()
    if err:
        return err
    data = request.json or {}
    cid = re.sub('[^a-z0-9-]', '-', (data.get('id') or '').lower().strip())
    if not cid:
        return (jsonify({'error': 'id required'}), 400)
    new_dir = _cur.CURRICULA_DIR / cid
    if new_dir.exists():
        return (jsonify({'error': 'already exists'}), 409)
    (new_dir / 'reference').mkdir(parents=True)
    (new_dir / 'papers').mkdir(parents=True)
    manifest = {'id': cid, 'name': data.get('name', cid), 'prompt_flavour': data.get('prompt_flavour', f"{data.get('name', cid)} exam preparation"), 'collection': cid.replace('-', '_'), 'reference_dir': f'curricula/{cid}/reference', 'papers_dir': f'curricula/{cid}/papers'}
    (new_dir / 'manifest.json').write_text(json.dumps(manifest, indent=2))
    (new_dir / 'topics.json').write_text('[]')
    return (jsonify(manifest), 201)

@app.route('/api/admin/curricula/<cid>/ingest', methods=['GET', 'POST'])
def api_admin_ingest(cid):
    err = _admin_only()
    if err:
        return err
    if request.method == 'GET':
        return jsonify(_INGEST_STATUS.get(cid, {'state': 'idle', 'chunks': 0}))
    try:
        manifest = _cur.load(cid)
    except FileNotFoundError:
        return (jsonify({'error': 'not found'}), 404)
    started = _start_ingest(cid, manifest)
    return jsonify({'ok': True, 'already_running': not started})

@app.route('/api/admin/curricula/<cid>/import-paper-questions', methods=['POST'])
def api_admin_import_paper_questions(cid):
    err = _admin_only()
    if err:
        return err
    import io
    from pypdf import PdfReader
    data = request.get_json(force=True) or {}
    url = (data.get('url') or '').strip()
    api_key = _key(data, 'api_key', 'API_KEY')
    model = (data.get('model') or 'gemma4').strip()
    if not url:
        return (jsonify({'error': 'url required'}), 400)
    if not _allowed_paper_host(url):
        return (jsonify({'error': 'Domain not in allowlist'}), 403)
    pdf_bytes = _fetch_pdf(url)
    if not pdf_bytes:
        return (jsonify({'error': 'Could not fetch PDF'}), 502)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    pages = [page.extract_text() or '' for page in reader.pages]
    full_text = '\n\n'.join(pages)
    if len(full_text.strip()) < 100:
        return (jsonify({'error': 'Could not extract text from PDF — may be scanned or maths-heavy'}), 422)
    extract_prompt = f'Output valid JSON only — no prose, no markdown fences.\n\nYou are given raw text from an exam past paper. Extract every question. Copy each question verbatim. If a question has numerical values or named variables (dates, quantities, percentages), change them to different plausible values so the question is fresh but structurally identical. Ignore cover pages, rubrics, instructions, and mark allocations. Return a JSON array. Each element: {{"question": "...", "answer": null}}. If a mark scheme is present in the text, populate answer.\n\nPAPER TEXT:\n{full_text[:14000]}'
    raw = generate(extract_prompt, model=model, api_key=api_key, provider='groq')
    m = re.search('\\[.*\\]', raw, re.DOTALL)
    if not m:
        return (jsonify({'error': 'LLM did not return a JSON array'}), 500)
    new_qs = json.loads(m.group())
    existing = json.loads(_questions_path(cid).read_text(encoding='utf-8')) if _questions_path(cid).exists() else []
    existing.extend(new_qs)
    _save_questions(cid, existing)
    return jsonify({'ok': True, 'imported': len(new_qs)})

@app.route('/api/admin/users')
def api_admin_users():
    err = _admin_only()
    if err:
        return err
    return jsonify(users.list_users())

@app.route('/api/admin/users/<int:uid>/role', methods=['PUT'])
def api_admin_set_role(uid: int):
    err = _admin_only()
    if err:
        return err
    role = (request.json or {}).get('role', '')
    if not role:
        return (jsonify({'error': 'role required'}), 400)
    try:
        users.set_role(uid, role)
    except users.UserError as e:
        return (jsonify({'error': str(e)}), 400)
    return jsonify({'ok': True})

@app.route('/api/admin/api-keys')
def api_admin_list_keys():
    err = _admin_only()
    if err:
        return err
    return jsonify(users.list_api_keys(g.user['id']))

@app.route('/api/admin/api-keys', methods=['POST'])
def api_admin_create_key():
    err = _admin_only()
    if err:
        return err
    label = ((request.json or {}).get('label') or '').strip()
    if not label:
        return (jsonify({'error': 'label required'}), 400)
    token = users.create_api_key(g.user['id'], label)
    return jsonify({'token': token})

@app.route('/api/admin/api-keys/<int:key_id>', methods=['DELETE'])
def api_admin_revoke_key(key_id: int):
    err = _admin_only()
    if err:
        return err
    users.revoke_api_key(key_id, g.user['id'])
    return jsonify({'ok': True})

def _dev_ok() -> bool:
    """True only if the request carries the DEV_PANEL_KEY. The panel is served at
    /dev/<key>; its API calls resend the key as an X-Dev-Key header. Constant-time,
    denies outright if the key is unset. Hides the whole /api/dev surface behind an
    unguessable path, independent of the auth_required flag."""
    import hmac
    expected = os.environ.get('DEV_PANEL_KEY', '')
    provided = request.headers.get('X-Dev-Key') or request.args.get('k', '')
    return bool(expected) and hmac.compare_digest(provided, expected)


@app.route('/dev/<key>')
def dev_panel(key):
    import hmac
    expected = os.environ.get('DEV_PANEL_KEY', '')
    if not (expected and hmac.compare_digest(key, expected)):
        return 'Not Found', 404
    err = _admin_only()
    if err:
        return err
    return send_from_directory('static', 'dev.html')

@app.route('/api/dev/settings')
def api_dev_get_settings():
    if not _dev_ok():
        return 'Not Found', 404
    err = _admin_only()
    if err:
        return err
    return jsonify(_settings.load())

@app.route('/api/dev/settings', methods=['PUT'])
def api_dev_put_settings():
    if not _dev_ok():
        return 'Not Found', 404
    err = _admin_only()
    if err:
        return err
    saved = _settings.save(request.get_json(force=True) or {})
    return jsonify(saved)
@app.route('/api/dev/sessions')
def api_dev_sessions():
    if not _dev_ok():
        return 'Not Found', 404
    err = _admin_only()
    if err:
        return err
    return jsonify({'sessions': all_sessions()})


@app.route('/api/dev/sessions/<sid>')
def api_dev_session_turns(sid):
    if not _dev_ok():
        return 'Not Found', 404
    err = _admin_only()
    if err:
        return err
    return jsonify({'turns': session_turns_admin(sid)})


@app.route('/api/dev/stats')
def api_dev_stats():
    if not _dev_ok():
        return 'Not Found', 404
    err = _admin_only()
    if err:
        return err
    return jsonify(stats())
if __name__ == '__main__':
    app.run(host='0.0.0.0', debug=True, port=5000, use_reloader=False)