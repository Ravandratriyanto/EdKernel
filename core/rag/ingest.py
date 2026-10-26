import glob
import json
import logging
import re
import warnings
from pathlib import Path
import chromadb
from chromadb.utils import embedding_functions
from pypdf import PdfReader
import pdftext
import io
warnings.filterwarnings('ignore')
logging.getLogger('pypdf').setLevel(logging.CRITICAL)
DB_PATH = 'data/chroma'
COLLECTION = 'reference'
_OK_PUNCT = '.,;:!?\'"()-%$/=+'
_MAX_JUNK = 0.08
_MIN_PAGE = 200
_MIN_CHUNK = 120

def _junk_ratio(s: str) -> float:
    if not s:
        return 1.0
    bad = sum((1 for c in s if not (c.isalnum() or c.isspace() or c in _OK_PUNCT)))
    return bad / len(s)

def _page_texts(pdf_path: str, vision_kwargs: dict | None=None):
    try:
        reader = PdfReader(pdf_path)
        n = len(reader.pages)
    except Exception:
        return
    pages = []
    for i in range(n):
        try:
            pages.append(pdftext.clean(pdftext.page_text(reader.pages[i])))
        except Exception:
            break
    pages = pdftext.strip_repeats(pages)
    if vision_kwargs:
        import base64
        import fitz
        from core.vision.image_handler import transcribe
        try:
            doc = fitz.open(pdf_path)
        except Exception:
            doc = None
        for i, text in enumerate(pages):
            if (not text.strip() or _junk_ratio(text) > _MAX_JUNK) and doc is not None:
                try:
                    pix = doc[i].get_pixmap(dpi=150)
                    b64 = base64.b64encode(pix.tobytes('jpeg')).decode()
                    pages[i] = transcribe(b64, name=f'p{i}.jpg', **vision_kwargs).strip()
                except Exception:
                    pass
    yield from enumerate(pages)

def _chunk(text: str, size: int=900, overlap: int=150) -> list[str]:
    text = re.sub('[ \\t]+', ' ', text)
    text = re.sub('\\n{3,}', '\n\n', text).strip()
    parts, buf = ([], '')
    for sentence in re.split('(?<=[.!?])\\s+', text):
        if len(buf) + len(sentence) + 1 > size and buf:
            parts.append(buf.strip())
            buf = buf[-overlap:] + ' ' + sentence
        else:
            buf += (' ' if buf else '') + sentence
    if buf.strip():
        parts.append(buf.strip())
    return [p for p in parts if len(p) > _MIN_CHUNK]

def _collect_pdfs(folder: str, vision_kwargs: dict | None=None) -> tuple[list, list, list]:
    docs, ids, metas = ([], [], [])
    for pdf in sorted(glob.glob(f'{folder}/*.pdf')):
        sidecar = Path(pdf.replace('.pdf', '.json'))
        title = 'unknown'
        if sidecar.exists():
            try:
                title = json.loads(sidecar.read_text()).get('title', 'unknown')
            except (ValueError, OSError):
                pass
        stem = Path(pdf).stem
        kept = 0
        for pageno, text in _page_texts(pdf, vision_kwargs=vision_kwargs):
            text = text.strip()
            if len(text) < _MIN_PAGE or _junk_ratio(text) > _MAX_JUNK:
                continue
            for k, chunk in enumerate(_chunk(text)):
                docs.append(chunk)
                ids.append(f'{stem}_p{pageno}_{k}')
                metas.append({'source': title[:80], 'page': pageno})
                kept += 1
        print(f'  {title[:52]:54} {kept:4} chunks')
    return (docs, ids, metas)

def pptx_text(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                parts.append(shape.text.strip())
    return '\n\n'.join(parts)

def _collect_texts(folder: str) -> tuple[list, list, list]:
    docs, ids, metas = ([], [], [])
    paths = sorted(p for p in Path(folder).glob('*') if p.suffix.lower() in ('.txt', '.pptx'))
    for path in paths:
        name = path.name
        if path.suffix.lower() == '.pptx':
            text = pptx_text(path.read_bytes())
        else:
            text = path.read_text(encoding='utf-8')
        chunks = _chunk(text)
        for k, chunk in enumerate(chunks):
            docs.append(chunk)
            ids.append(f'{name}_{k}')
            metas.append({'source': name, 'page': 0})
        print(f'  {name:54} {len(chunks):4} chunks')
    return (docs, ids, metas)

def build(text_folder: str='data/reference', pdf_folder: str='data/papers', db_path: str=DB_PATH, collection_name: str=COLLECTION, vision_kwargs: dict | None=None) -> int:
    client = chromadb.PersistentClient(path=db_path)
    ef = embedding_functions.DefaultEmbeddingFunction()
    try:
        client.delete_collection(collection_name)
    except Exception:
        pass
    collection = client.create_collection(name=collection_name, embedding_function=ef, metadata={'hnsw:space': 'cosine'})
    print('reference notes:')
    docs, ids, metas = _collect_texts(text_folder)
    print('past papers:')
    d2, i2, m2 = _collect_pdfs(pdf_folder, vision_kwargs=vision_kwargs)
    docs += d2
    ids += i2
    metas += m2
    if not docs:
        print('nothing to ingest')
        return 0
    for i in range(0, len(docs), 400):
        collection.upsert(documents=docs[i:i + 400], ids=ids[i:i + 400], metadatas=metas[i:i + 400])
        print(f'  embedded {min(i + 400, len(docs))}/{len(docs)}')
    print(f"\ncollection '{collection_name}': {collection.count()} chunks")
    return collection.count()
if __name__ == '__main__':
    import argparse
    import os
    import sys
    sys.path.insert(0, str(Path(__file__).parent.parent.parent))
    from dotenv import load_dotenv
    load_dotenv()
    p = argparse.ArgumentParser(description='Build a ChromaDB collection from reference material')
    p.add_argument('--curriculum', metavar='ID', help='Curriculum id (e.g. sat-alevel). Reads paths and collection name from curricula/<id>/manifest.json. If omitted, uses the legacy data/reference + data/papers paths.')
    p.add_argument('--vision', action='store_true', help='Use vision transcription for blank/maths-heavy PDF pages.')
    p.add_argument('--provider', default='ollama', help='LLM provider for vision (ollama or groq). Default: ollama.')
    p.add_argument('--model', default=None, help='Vision model name. Default: provider default.')
    p.add_argument('--host', default=None, help='Ollama host URL if not localhost.')
    args = p.parse_args()
    vkwargs = None
    if args.vision:
        vkwargs = {'provider': args.provider}
        if args.model:
            vkwargs['model'] = args.model
        if args.host:
            vkwargs['host'] = args.host
        api_key = os.getenv('API_KEY') or os.getenv('GROQ_API_KEY')
        if api_key:
            vkwargs['api_key'] = api_key
    if args.curriculum:
        import curricula as cur
        manifest = cur.load(args.curriculum)
        build(text_folder=cur.reference_dir(manifest), pdf_folder=cur.papers_dir(manifest), collection_name=manifest['collection'], vision_kwargs=vkwargs)
    else:
        build(vision_kwargs=vkwargs)